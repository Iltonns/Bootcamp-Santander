import openpyxl

wb = openpyxl.load_workbook('Ferramenta de Controle de Investimentos.xlsx')

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n{'='*60}")
    print(f"ABA: {sheet_name}")
    print(f"{'='*60}")
    
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=False):
        for cell in row:
            if cell.value is not None:
                if cell.data_type == 'f':
                    print(f"Célula {cell.coordinate}: Fórmula={cell.value}")
                else:
                    print(f"Célula {cell.coordinate}: {cell.value}")
